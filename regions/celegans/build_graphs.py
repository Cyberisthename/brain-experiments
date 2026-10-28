"""Build per-worm CSR graphs from Witvliet + Cook connectomes. Unsigned weights."""
from __future__ import annotations
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
import scipy.sparse as sp

H = Path(__file__).resolve().parent
RAW = H / "data" / "raw"
OUT = H / "data"
OUT.mkdir(parents=True, exist_ok=True)

WITV = [
    ("W1", "witvliet_2020_1_L1.xlsx", "L1", 1, "herm"),
    ("W2", "witvliet_2020_2_L1.xlsx", "L1", 1, "herm"),
    ("W3", "witvliet_2020_3_L1.xlsx", "L1", 1, "herm"),
    ("W4", "witvliet_2020_4_L1.xlsx", "L1", 1, "herm"),
    ("W5", "witvliet_2020_5_L2.xlsx", "L2", 2, "herm"),
    ("W6", "witvliet_2020_6_L3.xlsx", "L3", 3, "herm"),
    ("W7", "witvliet_2020_7_adult.xlsx", "adult", 4, "herm"),
    ("W8", "witvliet_2020_8_adult.xlsx", "adult", 4, "herm"),
]

HERM_NEURONS = set(
    p.strip()
    for p in (RAW / "PREFERRED_HERM_NEURON_NAMES.txt").read_text().splitlines()
    if p.strip()
)


def fix_name(n: str) -> str:
    import re
    n = str(n).strip()
    if n == "excgl":
        return "exc_gl"
    # VA01 -> VA1 to match OpenWorm NEURON_NAMES
    n = re.sub(r"([A-Za-z])0+([0-9])", r"\1\2", n)
    return n


def is_muscle_like(n: str) -> bool:
    if not n:
        return True
    if n[0].islower():
        return True
    prefixes = (
        "pm", "mu_", "vm", "dBWM", "vBWM", "BWM", "anal", "intL", "intR",
        "hyp", "glr", "cepsh", "GLR",
    )
    # pharyngeal muscles pm*; body wall; exclude neuron MCL/MCR/MI
    if n in ("MCL", "MCR", "MI", "M1", "M2L", "M2R", "M3L", "M3R", "M4", "M5"):
        return False
    if n.startswith(prefixes):
        return True
    if n.startswith("mc") and n not in ("MCL", "MCR"):
        # mc2DL etc muscles; MCL is neuron
        if len(n) > 2 and n[2].isdigit():
            return True
    return False


def edges_to_csr(nodes, edge_w):
    """edge_w: dict (pre, post) -> weight; W[post, pre] = w."""
    names = sorted(nodes)
    idx = {n: i for i, n in enumerate(names)}
    rows, cols, data = [], [], []
    for (pre, post), w in edge_w.items():
        if pre not in idx or post not in idx:
            continue
        if w == 0:
            continue
        rows.append(idx[post])
        cols.append(idx[pre])
        data.append(float(w))
    N = len(names)
    W = sp.csr_matrix((data, (rows, cols)), shape=(N, N), dtype=np.float64)
    W.eliminate_zeros()
    return W, np.array(names)


def save_worm(worm_id, W, names, meta):
    sign = np.ones(len(names), dtype=np.int8)
    path = OUT / f"{worm_id}.npz"
    np.savez_compressed(
        path,
        W_data=W.data,
        W_indices=W.indices,
        W_indptr=W.indptr,
        W_shape=np.array(W.shape),
        names=names,
        sign=sign,
        ei_status=np.array("unknown"),
        meta_json=np.array(json.dumps(meta)),
    )
    return {
        "id": worm_id,
        "N": int(W.shape[0]),
        "nnz": int(W.nnz),
        "path": str(path.name),
        **{k: meta[k] for k in ("stage", "stage_ord", "sex", "source", "n_chem", "n_elec")},
    }


def load_witvliet(worm_id, fname, stage, stage_ord, sex):
    path = RAW / fname
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    edge_w = defaultdict(float)
    n_chem = n_elec = 0
    nodes = set()
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or row[0] is None:
            continue
        pre, post, typ, syn = row[0], row[1], row[2], row[3]
        pre, post = fix_name(pre), fix_name(post)
        # keep neuronal cells; drop obvious non-neurons
        if is_muscle_like(pre) or is_muscle_like(post):
            continue
        w = float(syn or 0)
        if w <= 0:
            continue
        nodes.add(pre)
        nodes.add(post)
        typ = str(typ).lower()
        if typ.startswith("chem"):
            edge_w[(pre, post)] += w
            n_chem += 1
        elif typ.startswith("elec"):
            # symmetric
            edge_w[(pre, post)] += w
            edge_w[(post, pre)] += w
            n_elec += 1
        else:
            edge_w[(pre, post)] += w
    wb.close()
    W, names = edges_to_csr(nodes, edge_w)
    meta = dict(
        stage=stage,
        stage_ord=stage_ord,
        sex=sex,
        source="Witvliet2021",
        file=fname,
        n_chem=n_chem,
        n_elec=n_elec,
        ei_status="unknown",
        gap_note="Witvliet electrical annotations incomplete",
    )
    return save_worm(worm_id, W, names, meta)


def load_cook_edgelist_herm():
    """Cook herm from OpenWorm herm_full_edgelist.csv, neurons only."""
    import csv

    path = RAW / "herm_full_edgelist.csv"
    edge_w = defaultdict(float)
    n_chem = n_elec = 0
    nodes = set()
    with open(path, newline="") as f:
        r = csv.DictReader(f)
        for row in r:
            pre = fix_name(row["Source"])
            post = fix_name(row["Target"])
            typ = str(row["Type"]).strip().lower()
            w = float(row["Weight"])
            if pre not in HERM_NEURONS or post not in HERM_NEURONS:
                continue
            if w <= 0:
                continue
            nodes.add(pre)
            nodes.add(post)
            if typ.startswith("chem"):
                edge_w[(pre, post)] += w
                n_chem += 1
            elif typ.startswith("elec"):
                edge_w[(pre, post)] += w
                edge_w[(post, pre)] += w
                n_elec += 1
    W, names = edges_to_csr(nodes, edge_w)
    meta = dict(
        stage="adult",
        stage_ord=4,
        sex="herm",
        source="Cook2019",
        file="herm_full_edgelist.csv",
        n_chem=n_chem,
        n_elec=n_elec,
        ei_status="unknown",
        weight_note="EM sections (Cook), not raw synapse counts",
    )
    return save_worm("COOK_H", W, names, meta)


def load_cook_si5_male():
    """Parse Cook SI5 male chemical + gap symmetric; drop muscle-like cells."""
    path = RAW / "Cook2019_SI5_corrected.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    edge_w = defaultdict(float)
    n_chem = n_elec = 0

    def parse_sheet(sheet_name, electrical=False):
        nonlocal n_chem, n_elec
        ws = wb[sheet_name]
        # header names in row 3, cols from 4; pre names in col C from row 4
        max_c = ws.max_column
        max_r = ws.max_row
        posts = []
        for c in range(4, max_c + 1):
            v = ws.cell(3, c).value
            posts.append(fix_name(v) if v is not None else None)
        for r in range(4, max_r + 1):
            pre = ws.cell(r, 3).value
            if pre is None:
                continue
            pre = fix_name(pre)
            if is_muscle_like(pre):
                continue
            for j, post in enumerate(posts):
                if post is None or is_muscle_like(post):
                    continue
                val = ws.cell(r, 4 + j).value
                if val is None:
                    continue
                w = float(val)
                if w <= 0:
                    continue
                if electrical:
                    edge_w[(pre, post)] += w
                    # symmetric sheet already has both; still OK to add as listed
                    n_elec += 1
                else:
                    edge_w[(pre, post)] += w
                    n_chem += 1

    parse_sheet("male chemical", electrical=False)
    parse_sheet("male gap jn symmetric", electrical=True)
    wb.close()
    nodes = set()
    for pre, post in edge_w:
        nodes.add(pre)
        nodes.add(post)
    # drop edges to/from muscle that slipped in
    nodes = {n for n in nodes if not is_muscle_like(n)}
    W, names = edges_to_csr(nodes, edge_w)
    meta = dict(
        stage="adult",
        stage_ord=4,
        sex="male",
        source="Cook2019",
        file="Cook2019_SI5_corrected.xlsx",
        n_chem=n_chem,
        n_elec=n_elec,
        ei_status="unknown",
        weight_note="EM sections (Cook); muscle cells excluded by heuristic",
    )
    return save_worm("COOK_M", W, names, meta)


def main():
    manifest = []
    for wid, fn, stage, so, sex in WITV:
        info = load_witvliet(wid, fn, stage, so, sex)
        print(info)
        manifest.append(info)
    info = load_cook_edgelist_herm()
    print(info)
    manifest.append(info)
    info = load_cook_si5_male()
    print(info)
    manifest.append(info)
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2))
    ok = sum(1 for m in manifest if m["id"].startswith("W") and m["N"] >= 100 and m["nnz"] >= 200)
    print(f"GATE witvliet_ok={ok}/8")
    if ok < 6:
        raise SystemExit("GATE FAIL: fewer than 6 Witvliet worms usable")
    print("build_graphs done")


if __name__ == "__main__":
    main()
